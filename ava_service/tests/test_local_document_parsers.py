from io import BytesIO
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from core.enums import InputModality
from preprocessing.multimodal import (
    MultimodalPreprocessor,
    MultimodalRequest,
    PreprocessingError,
    RawAttachment,
    parse_local_document,
)


def test_txt_and_csv_are_extracted_with_utf8_bom_support():
    txt = RawAttachment(file_name="note.txt", content_bytes=b"\xef\xbb\xbfXin chao")
    csv_file = RawAttachment(
        file_name="attendance.csv",
        content_bytes="Tên,Ngày\nAn,2026-09-20\n".encode(),
    )
    assert parse_local_document(txt) == "Xin chao"
    assert "Tên | Ngày" in parse_local_document(csv_file)
    assert "An | 2026-09-20" in parse_local_document(csv_file)


def test_docx_parser_preserves_table_boundaries():
    document = Document()
    document.add_paragraph("Chính sách nghỉ phép")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nhân viên"
    table.cell(0, 1).text = "12 ngày"
    stream = BytesIO()
    document.save(stream)

    extracted = parse_local_document(
        RawAttachment(file_name="policy.docx", content_bytes=stream.getvalue())
    )
    assert "Chính sách nghỉ phép" in extracted
    assert "[TABLE 1]" in extracted
    assert "Nhân viên | 12 ngày" in extracted


def test_docx_parser_reads_text_inside_textboxes_and_content_controls():
    document = Document()
    document.add_paragraph("Nội dung thường")
    stream = BytesIO()
    document.save(stream)

    source = BytesIO(stream.getvalue())
    result = BytesIO()
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(result, "w") as outgoing:
        for item in incoming.infolist():
            value = incoming.read(item.filename)
            if item.filename == "word/document.xml":
                xml = value.decode("utf-8")
                injected = (
                    '<w:sdt><w:sdtContent><w:p><w:r><w:t>'
                    'Lệnh HRMS nằm trong hộp nội dung'
                    '</w:t></w:r></w:p></w:sdtContent></w:sdt>'
                )
                xml = xml.replace("</w:body>", injected + "</w:body>")
                value = xml.encode("utf-8")
            outgoing.writestr(item, value)

    extracted = parse_local_document(
        RawAttachment(file_name="commands.docx", content_bytes=result.getvalue())
    )

    assert "Nội dung thường" in extracted
    assert "Lệnh HRMS nằm trong hộp nội dung" in extracted


def test_utf16_text_document_is_supported():
    attachment = RawAttachment(
        file_name="vietnamese.txt",
        content_bytes="Nội dung tiếng Việt".encode("utf-16"),
    )

    assert parse_local_document(attachment) == "Nội dung tiếng Việt"


def test_xlsx_parser_preserves_worksheet_name_and_values():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ChamCong"
    sheet.append(["Nhân viên", "Ngày"])
    sheet.append(["An", "2026-09-20"])
    stream = BytesIO()
    workbook.save(stream)

    extracted = parse_local_document(
        RawAttachment(
            file_name="attendance.xlsx",
            content_bytes=stream.getvalue(),
        )
    )
    assert "[SHEET ChamCong]" in extracted
    assert "An | 2026-09-20" in extracted


def test_scanned_pdf_without_local_ocr_fails_explicitly():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(PreprocessingError, match="OCR_UNAVAILABLE"):
        parse_local_document(
            RawAttachment(file_name="scan.pdf", content_bytes=stream.getvalue())
        )


def test_image_without_local_ocr_or_vision_fails_explicitly():
    request = MultimodalRequest(
        text="Đọc ảnh này",
        modality=InputModality.IMAGE,
        attachments=[
            RawAttachment(file_name="scan.png", content_bytes=b"not-a-real-png")
        ],
    )
    with pytest.raises(PreprocessingError, match="IMAGE_INVALID"):
        MultimodalPreprocessor().preprocess(request)


def _png_bytes() -> bytes:
    from PIL import Image
    stream = BytesIO()
    Image.new("RGB", (32, 32), "white").save(stream, format="PNG")
    return stream.getvalue()


def test_image_ocr_and_vision_are_separate_adapters():
    request = MultimodalRequest(
        text="Lỗi này là gì?",
        modality=InputModality.IMAGE,
        attachments=[RawAttachment(file_name="screen.png", content_bytes=_png_bytes())],
    )
    normalized = MultimodalPreprocessor(
        image_ocr=lambda _: "Mã lỗi HR-404",
        image_vision=lambda _: '{"description":"Trang báo lỗi","elements":["alert"]}',
    ).preprocess(request)
    assert normalized.attachments[0].extracted_text == "Mã lỗi HR-404"
    assert "visual_observation" in normalized.attachments[0].metadata


def test_vision_failure_falls_back_to_successful_ocr():
    request = MultimodalRequest(
        text="Mô tả ảnh",
        modality=InputModality.IMAGE,
        attachments=[
            RawAttachment(
                file_name="screen.png",
                content_bytes=_png_bytes(),
            )
        ],
    )

    def unavailable(_):
        raise PreprocessingError("VISION_UNAVAILABLE")

    normalized = MultimodalPreprocessor(
        image_ocr=lambda _: "visible text",
        image_vision=unavailable,
    ).preprocess(request)

    assert normalized.attachments[0].extracted_text == "visible text"


def test_ocr_failure_falls_back_to_successful_vision():
    request = MultimodalRequest(
        text="Mô tả ảnh",
        modality=InputModality.IMAGE,
        attachments=[
            RawAttachment(
                file_name="screen.png",
                content_bytes=_png_bytes(),
            )
        ],
    )

    def unavailable(_):
        raise PreprocessingError("OCR_UNAVAILABLE")

    normalized = MultimodalPreprocessor(
        image_ocr=unavailable,
        image_vision=lambda _: (
            '{"description":"Form",'
            '"visibleText":["NV002","Status OK"]}'
        ),
    ).preprocess(request)

    visual = normalized.attachments[
        0
    ].metadata.get("visual_observation")

    assert visual
    assert "NV002" in visual


def test_image_fails_only_when_ocr_and_vision_both_fail():
    request = MultimodalRequest(
        text="Mô tả ảnh",
        modality=InputModality.IMAGE,
        attachments=[
            RawAttachment(
                file_name="screen.png",
                content_bytes=_png_bytes(),
            )
        ],
    )

    def ocr_unavailable(_):
        raise PreprocessingError("OCR_UNAVAILABLE")

    def vision_unavailable(_):
        raise PreprocessingError("VISION_UNAVAILABLE")

    with pytest.raises(
        PreprocessingError,
        match="IMAGE_ANALYSIS_UNAVAILABLE",
    ):
        MultimodalPreprocessor(
            image_ocr=ocr_unavailable,
            image_vision=vision_unavailable,
        ).preprocess(request)


def test_invalid_voice_fails_explicitly():
    request = MultimodalRequest(
        text="",
        modality=InputModality.VOICE,
        attachments=[RawAttachment(file_name="voice.wav", content_bytes=b"RIFF\x24\x00\x00\x00WAVEfmt ")],
    )
    with pytest.raises(PreprocessingError, match="AUDIO_INVALID"):
        MultimodalPreprocessor().preprocess(request)


def test_document_parser_rejects_missing_bytes():
    with pytest.raises(PreprocessingError, match="ATTACHMENT_BYTES_REQUIRED"):
        parse_local_document(RawAttachment(file_name="note.txt"))
